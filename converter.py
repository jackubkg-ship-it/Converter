import io
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

THIN = Side(style='thin', color='999999')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

TITLE_FONT = Font(name='Calibri', size=12, bold=True, color='1F3864')
HEADER_FONT = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
HEADER_FILL = PatternFill('solid', fgColor='305496')
CELL_FONT = Font(name='Calibri', size=10)
BOLD_FONT = Font(name='Calibri', size=10, bold=True)

HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
CELL_ALIGN = Alignment(vertical='top', wrap_text=True)
NUM_ALIGN = Alignment(horizontal='right', vertical='top')
TOTAL_FILL = PatternFill('solid', fgColor='D9E1F2')

COL_WIDTHS_MAIN = [8, 22, 16, 10, 12, 14, 14, 12, 16, 20, 30, 22, 22, 30, 20, 12, 10, 10, 14]
COL_WIDTHS_OTHER = [8, 22, 16, 10, 16, 12, 14, 14, 12, 16, 20, 30, 22, 22, 30, 20, 12, 10, 10, 14]


def parse_date(val):
    if isinstance(val, datetime):
        return val
    if isinstance(val, str):
        for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
            try:
                return datetime.strptime(val, fmt)
            except ValueError:
                pass
    return None


def _fmt_num(n):
    if n is None or n == "":
        return "0"
    if isinstance(n, float) and n.is_integer():
        return str(int(n))
    if isinstance(n, int):
        return str(n)
    return str(n)


def format_cost(qty, price, total):
    return f"{_fmt_num(qty)}x{_fmt_num(price)}={_fmt_num(total)}"


def is_empty_date(d):
    if d is None:
        return True
    if isinstance(d, str) and d.strip() == "":
        return True
    return False


def apply_column_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_title_block(ws, last_col_letter):
    ws.merge_cells(f'A1:{last_col_letter}1')
    ws['A1'] = '"Salyan Oil" Ltd şirkəti'
    ws['A1'].font = TITLE_FONT
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')

    ws.merge_cells(f'A2:{last_col_letter}2')
    ws['A2'] = '"Avto 326" firmasından'
    ws['A2'].font = TITLE_FONT
    ws['A2'].alignment = Alignment(horizontal='center', vertical='center')

    ws.merge_cells(f'A3:{last_col_letter}3')
    ws['A3'] = "Avtomobillərə göstərilən təmirin hesabatı"
    ws['A3'].font = TITLE_FONT
    ws['A3'].alignment = Alignment(horizontal='center', vertical='center')

    ws.row_dimensions[1].height = 20
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 20


def write_header_row(ws, headers, row=4):
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = HEADER_ALIGN
        c.border = BORDER
    ws.row_dimensions[row].height = 42


def style_data_row(ws, row, last_col):
    for col in range(1, last_col + 1):
        c = ws.cell(row=row, column=col)
        c.font = CELL_FONT
        c.border = BORDER
        c.alignment = CELL_ALIGN


def convert_workbook(input_stream, split_sheets=True, add_totals=True):
    wb_in = load_workbook(input_stream, data_only=True)

    sheet_name = None
    for name in wb_in.sheetnames:
        if "İyul2026" in name or "İyul" in name:
            sheet_name = name
            break
    if sheet_name is None:
        sheet_name = wb_in.sheetnames[0]
    ws_in = wb_in[sheet_name]

    wb_out = Workbook()
    wb_out.remove(wb_out.active)

    ws_main = wb_out.create_sheet("Avqust 2026")
    ws_batt = wb_out.create_sheet("Akumlyator")

    header_main = [
        "Sıra nömrəsi", "Avtomobilin Markası", "Dovlət qeydiyyat nişanı", "Buraxiliş ili",
        "Servisa daxil olan tarix", "Servisdan planlaşdırılan çıxma tarixi", "Servisdan faktiki çıxma tarixi",
        "Spidometr göstəricisi", "Sürücünün adı", "Avtomobili təmirə sorğunu göndərən şəxsin adı",
        "Dəyişdirilən ehtiyyat hissəsinin adı", "Ehtiyyat hissəsinin qiyməti (AZN)",
        "Görülən işin qiyməti (AZN)", "Təmirin qısa təsviri",
        "Sahədə təmiri təsdiqləyən şəxsin/lərin adı", "Cəmi", "0.15", "0.2", "Toplam xərc"
    ]
    header_batt = [
        "Sıra nömrəsi", "Avtomobilin Markası", "Dovlət qeydiyyat nişanı", "Buraxiliş ili",
        "Podaratçı şirkət", "Servisa daxil olan tarix", "Servisdan planlaşdırılan çıxma tarixi",
        "Servisdan faktiki çıxma tarixi", "Spidometr göstəricisi", "Sürücünün adı",
        "Avtomobili təmirə sorğunu göndərən şəxsin adı", "Dəyişdirilən ehtiyyat hissəsinin adı",
        "Ehtiyyat hissəsinin qiyməti (AZN)", "Görülən işin qiyməti (AZN)",
        "Təmirin qısa təsviri", "Sahədə təmiri təsdiqləyən şəxsin/lərin adı",
        "Cəmi", "0.15", "0.2", "Toplam xərc"
    ]

    write_title_block(ws_main, 'S')
    write_title_block(ws_batt, 'T')

    write_header_row(ws_main, header_main, row=4)
    write_header_row(ws_batt, header_batt, row=4)

    apply_column_widths(ws_main, COL_WIDTHS_MAIN)
    apply_column_widths(ws_batt, COL_WIDTHS_OTHER)

    ws_main.freeze_panes = 'A5'
    ws_batt.freeze_panes = 'A5'

    groups = {}
    group_order = []
    batt_rows = []
    last_group_by_vehicle = {}

    for row in ws_in.iter_rows(min_row=3, values_only=False):
        if all(cell.value is None for cell in row):
            continue

        tarix = row[1].value
        desc = row[2].value
        say = row[3].value
        unit = row[4].value
        qiymet = row[5].value
        yekun = row[6].value
        dq = row[12].value
        nv = row[13].value

        if not desc and not nv:
            continue

        desc_str = str(desc).strip() if desc else ""
        unit_str = str(unit).strip().lower() if unit else ""
        nv_str = str(nv).strip() if nv else ""
        dq_str = str(dq).strip() if dq else ""

        if yekun is None and say is not None and qiymet is not None:
            try:
                yekun = float(say) * float(qiymet)
            except Exception:
                yekun = 0.0
        else:
            try:
                yekun = float(yekun) if yekun is not None else 0.0
            except Exception:
                yekun = 0.0

        date_val = parse_date(tarix) or tarix
        date_missing = is_empty_date(date_val)
        desc_lower = desc_str.lower()
        is_work = unit_str in ("iş", "is")

        # ---------- 0.15 / 0.2 ----------
        if desc_str in ("0.15", "0.2"):
            if date_missing:
                veh_key = (nv_str, dq_str)
                if veh_key in last_group_by_vehicle:
                    key = last_group_by_vehicle[veh_key]
                else:
                    key = (nv_str, dq_str, date_val)
                    if key not in groups:
                        groups[key] = dict(nv=nv_str, dq=dq_str, date=date_val,
                                           items=[], tax_015=0.0, tax_02=0.0)
                        group_order.append(key)
                    last_group_by_vehicle[veh_key] = key
            else:
                key = (nv_str, dq_str, date_val)
                if key not in groups:
                    groups[key] = dict(nv=nv_str, dq=dq_str, date=date_val,
                                       items=[], tax_015=0.0, tax_02=0.0)
                    group_order.append(key)
                last_group_by_vehicle[(nv_str, dq_str)] = key

            if desc_str == "0.15":
                groups[key]["tax_015"] += yekun
            else:
                groups[key]["tax_02"] += yekun
            continue

        # ---------- Аккумуляторы — только они идут на отдельный лист ----------
        if not is_work and ("akkumulyator" in desc_lower or "akkum" in desc_lower):
            batt_rows.append(dict(nv=nv_str, dq=dq_str, date=date_val,
                                  desc=desc_str, say=say, qiymet=qiymet, yekun=yekun))
            continue

        # ---------- Всё остальное (включая масла) — в основной лист ----------
        key = (nv_str, dq_str, date_val)
        if key not in groups:
            groups[key] = dict(nv=nv_str, dq=dq_str, date=date_val,
                               items=[], tax_015=0.0, tax_02=0.0)
            group_order.append(key)
        last_group_by_vehicle[(nv_str, dq_str)] = key

        groups[key]["items"].append(dict(
            desc=desc_str, say=say, qiymet=qiymet, yekun=yekun, is_work=is_work
        ))

    # ---------- Основной лист ----------
    row_idx = 5
    for i, key in enumerate(group_order, 1):
        g = groups[key]
        ws_main.cell(row=row_idx, column=1, value=i)
        ws_main.cell(row=row_idx, column=2, value=g["nv"])
        ws_main.cell(row=row_idx, column=3, value=g["dq"])
        d = g["date"]
        if isinstance(d, datetime):
            ws_main.cell(row=row_idx, column=5, value=d.strftime("%d.%m.%Y"))
        else:
            ws_main.cell(row=row_idx, column=5, value=str(d) if d else "")

        part_descs, work_descs = [], []
        part_costs, work_costs = [], []
        total_parts = 0.0
        total_works = 0.0

        for it in g["items"]:
            if it["is_work"]:
                work_descs.append(it["desc"])
                work_costs.append(format_cost(it["say"], it["qiymet"], it["yekun"]))
                total_works += it["yekun"]
            else:
                part_descs.append(it["desc"])
                part_costs.append(format_cost(it["say"], it["qiymet"], it["yekun"]))
                total_parts += it["yekun"]

        all_descs = part_descs + work_descs
        ws_main.cell(row=row_idx, column=11, value="\n".join(all_descs))
        ws_main.cell(row=row_idx, column=12, value="\n".join(part_costs))
        ws_main.cell(row=row_idx, column=13, value="\n".join(work_costs))

        total = total_parts + total_works
        c16 = ws_main.cell(row=row_idx, column=16, value=total)
        c16.number_format = '#,##0.00'
        c17 = ws_main.cell(row=row_idx, column=17, value=g["tax_015"])
        c17.number_format = '#,##0.00'
        c18 = ws_main.cell(row=row_idx, column=18, value=g["tax_02"])
        c18.number_format = '#,##0.00'
        c19 = ws_main.cell(row=row_idx, column=19,
                           value=f"=P{row_idx}+Q{row_idx}+R{row_idx}")
        c19.number_format = '#,##0.00'

        style_data_row(ws_main, row_idx, 19)
        c19.font = BOLD_FONT
        row_idx += 1

    # ---------- Лист Akumlyator ----------
    row_idx = 5
    for i, b in enumerate(batt_rows, 1):
        ws_batt.cell(row=row_idx, column=1, value=i)
        ws_batt.cell(row=row_idx, column=2, value=b["nv"])
        ws_batt.cell(row=row_idx, column=3, value=b["dq"])
        ws_batt.cell(row=row_idx, column=5, value="Avto 326")
        d = b["date"]
        if isinstance(d, datetime):
            ws_batt.cell(row=row_idx, column=6, value=d.strftime("%d.%m.%Y"))
        else:
            ws_batt.cell(row=row_idx, column=6, value=str(d) if d else "")
        ws_batt.cell(row=row_idx, column=12, value=b["desc"])
        ws_batt.cell(row=row_idx, column=13,
                     value=format_cost(b["say"], b["qiymet"], b["yekun"]))
        c17 = ws_batt.cell(row=row_idx, column=17, value=b["yekun"])
        c17.number_format = '#,##0.00'
        c20 = ws_batt.cell(row=row_idx, column=20,
                           value=f"=Q{row_idx}+R{row_idx}+S{row_idx}")
        c20.number_format = '#,##0.00'
        style_data_row(ws_batt, row_idx, 20)
        c20.font = BOLD_FONT
        row_idx += 1

    # ---------- Итоги ----------
    if add_totals:
        for ws, start_col, end_col, label_col in [
            (ws_main, 16, 19, 16),
            (ws_batt, 17, 20, 17),
        ]:
            last_row = ws.max_row
            if last_row >= 5:
                tr = last_row + 1
                lbl = ws.cell(row=tr, column=label_col, value="Cəmi")
                lbl.font = BOLD_FONT
                lbl.fill = TOTAL_FILL
                lbl.alignment = Alignment(horizontal='right', vertical='center')
                lbl.border = BORDER
                for col in range(start_col, end_col + 1):
                    L = get_column_letter(col)
                    c = ws.cell(row=tr, column=col,
                                value=f"=SUM({L}5:{L}{last_row})")
                    c.font = BOLD_FONT
                    c.fill = TOTAL_FILL
                    c.number_format = '#,##0.00'
                    c.alignment = NUM_ALIGN
                    c.border = BORDER

    output = io.BytesIO()
    wb_out.save(output)
    output.seek(0)
    return output